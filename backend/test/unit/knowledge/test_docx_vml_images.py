"""真实 DOCX 的 VML 空图形分类与图片完整性回归。"""

import io
import re

import pytest
from docx import Document
from lxml import etree
from PIL import Image

from yuxi.infrastructure.document_parsing.parser import parse

VML = "urn:schemas-microsoft-com:vml"
WORD = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
OFFICE = "urn:schemas-microsoft-com:office:office"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def append_vml(document, attributes, *, content=None, child=False):
    """在真实 Word 段落中加入明确构造的 VML 矩形。"""
    pict = etree.Element(f"{{{WORD}}}pict")
    rect = etree.SubElement(pict, f"{{{VML}}}rect", filled="t", stroked="f")
    image = etree.SubElement(rect, f"{{{VML}}}imagedata", attributes)
    image.text = content
    if child:
        etree.SubElement(image, f"{{{VML}}}unexpected")
    document.add_paragraph().add_run()._r.append(pict)


def image_bytes():
    """生成尺寸与颜色确定的图片，避免被当成透明间隔图。"""
    stream = io.BytesIO()
    Image.new("RGB", (32, 24), (15, 90, 180)).save(stream, format="PNG")
    return stream.getvalue()


@pytest.mark.asyncio
@pytest.mark.parametrize("attributes", [{}, {f"{{{OFFICE}}}title": ""}])
async def test_empty_vml_preserves_text_image_position_and_source(tmp_path, attributes):
    """空 VML 不阻止有效图片；正文顺序、像素和原件保持完整。"""
    document = Document()
    document.add_paragraph("before")
    append_vml(document, attributes)
    document.add_paragraph("middle")
    document.add_picture(io.BytesIO(image_bytes()))
    append_vml(document, attributes)
    document.add_paragraph("after")
    source = tmp_path / "source.docx"
    document.save(source)
    original = source.read_bytes()

    result = await parse(source, tmp_path / "parsed")

    markdown = result.markdown_path.read_text()
    assert re.search(r"before[\s\S]+middle[\s\S]+!\[image\]\(images/figure-001.png\)[\s\S]+after", markdown)
    assert "<!-- image -->" not in markdown
    assert result.resources == ("images/figure-001.png",)
    with Image.open(result.directory / result.resources[0]) as image:
        assert image.size == (32, 24)
        assert image.convert("RGB").tobytes() == Image.new("RGB", (32, 24), (15, 90, 180)).tobytes()
    assert source.read_bytes() == original


@pytest.mark.asyncio
async def test_vml_with_valid_image_relationship_preserves_resource(tmp_path):
    """具有真实关系的 VML 图片不能被归入空图形。"""
    document = Document()
    document.add_paragraph("before")
    relationship, _ = document.part.get_or_add_image(io.BytesIO(image_bytes()))
    append_vml(document, {f"{{{REL}}}id": relationship})
    document.add_paragraph("after")
    source = tmp_path / "source.docx"
    document.save(source)

    result = await parse(source, tmp_path / "parsed")

    markdown = result.markdown_path.read_text()
    assert re.search(r"before[\s\S]+!\[image\]\(images/figure-001.png\)[\s\S]+after", markdown)
    assert result.resources == ("images/figure-001.png",)
    with Image.open(result.directory / result.resources[0]) as image:
        assert image.convert("RGB").tobytes() == Image.new("RGB", (32, 24), (15, 90, 180)).tobytes()


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("attributes", "content", "child"),
    [
        ({f"{{{REL}}}id": "missing"}, None, False),
        ({f"{{{REL}}}id": ""}, None, False),
        ({f"{{{OFFICE}}}relid": "missing"}, None, False),
        ({"src": "missing.png"}, None, False),
        ({"unknown": "value"}, None, False),
        ({}, "unexpected content", False),
        ({}, None, True),
    ],
)
async def test_referenced_or_unknown_vml_without_image_rejects_partial_result(tmp_path, attributes, content, child):
    """真实缺图或无法确认的节点仍失败，不能发布部分 Markdown。"""
    document = Document()
    document.add_paragraph("before")
    document.add_picture(io.BytesIO(image_bytes()))
    append_vml(document, attributes, content=content, child=child)
    document.add_paragraph("after")
    source = tmp_path / "source.docx"
    document.save(source)
    original = source.read_bytes()

    with pytest.raises(ValueError, match="Office 图片缺少图像数据"):
        await parse(source, tmp_path / "parsed")

    assert not (tmp_path / "parsed").exists()
    assert source.read_bytes() == original
